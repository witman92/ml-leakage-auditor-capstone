# ==============================================================================
# build_presentation.py · Genera la presentación de la entrega final de Capstone I
# ==============================================================================
"""
Construye 07_report/presentation/capstone_presentation_ef.pptx a partir de la plantilla visual UDLA
(assets/udla_base_template.pptx): conserva la portada, el maestro (banda naranja y logo) y la
diapositiva de cierre, y reemplaza todo el contenido por el del proyecto vigente.

Las figuras se leen de 04_experiments/results/ (las mismas del informe) y cada diapositiva lleva
en sus notas el guion del orador con las cifras y las preguntas probables del comité.

    python 07_report/presentation/build_presentation.py
"""

# %% PASO 0 · Importaciones, rutas y paleta de la plantilla
# ------------------------------------------------------------------------------
# La paleta y las tipografías se extrajeron de la plantilla base para mantener el mismo estilo.
# ------------------------------------------------------------------------------
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
PLANTILLA = AQUI / "assets" / "udla_base_template.pptx"
SALIDA = AQUI / "capstone_presentation_ef.pptx"
FIGURAS = RAIZ / "04_experiments" / "results"

TINTA = RGBColor(0x2D, 0x37, 0x48)      # títulos y texto
AZUL = RGBColor(0x4A, 0x7A, 0xA7)
ROJO = RGBColor(0xE0, 0x5A, 0x47)
AMBAR = RGBColor(0xE8, 0xA3, 0x17)
VERDE = RGBColor(0x3C, 0x9A, 0x5F)
NARANJO = RGBColor(0xEE, 0x6A, 0x1F)
GRIS_PANEL = RGBColor(0xF5, 0xF6, 0xF8)
GRIS_BORDE = RGBColor(0xDD, 0xE1, 0xE6)
GRIS_FLECHA = RGBColor(0x5A, 0x64, 0x72)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
FUENTE = "Calibri"


# %% PASO 1 · Utilidades de dibujo (mismos componentes que la plantilla)
# ------------------------------------------------------------------------------
# texto(): cuadro de texto; tarjeta(): caja blanca con barra superior de color; panel(): caja gris
# con barra lateral; dato(): cifra grande con etiqueta; mensaje(): franja azul inferior.
# ------------------------------------------------------------------------------
def texto(s, x, y, w, h, contenido, size=16, color=TINTA, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, fuente=FUENTE):
    """contenido: str o lista de líneas; una línea puede ser (texto, {opciones}) para variar el formato."""
    caja = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = caja.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lineas = contenido if isinstance(contenido, list) else [contenido]
    for i, linea in enumerate(lineas):
        t, op = (linea, {}) if isinstance(linea, str) else linea
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(op.get("after", 4))
        r = p.add_run()
        r.text = t
        f = r.font
        f.name = op.get("fuente", fuente)
        f.size = Pt(op.get("size", size))
        f.bold = op.get("bold", bold)
        f.color.rgb = op.get("color", color)
    return caja


def caja(s, x, y, w, h, relleno, borde=None, forma=MSO_SHAPE.ROUNDED_RECTANGLE):
    c = s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = relleno
    if borde is None:
        c.line.fill.background()
    else:
        c.line.color.rgb = borde
        c.line.width = Pt(1)
    c.shadow.inherit = False
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        c.adjustments[0] = 0.05
    return c


def titulo(s, t):
    texto(s, 0.6, 0.82, 12.1, 0.7, t, size=30, bold=True)


def tarjeta(s, x, y, w, h, color, cabeza, cuerpo, size=15):
    caja(s, x, y, w, h, BLANCO, GRIS_BORDE)
    caja(s, x, y, w, 0.09, color, forma=MSO_SHAPE.RECTANGLE)
    texto(s, x + 0.15, y + 0.18, w - 0.3, 0.5, cabeza, size=18, bold=True, color=color)
    texto(s, x + 0.15, y + 0.72, w - 0.3, h - 0.8, cuerpo, size=size)


def panel(s, x, y, w, h, color, cabeza, cuerpo, size=13.5, size_cabeza=18):
    caja(s, x, y, w, h, GRIS_PANEL)
    caja(s, x, y, 0.09, h, color, forma=MSO_SHAPE.RECTANGLE)
    texto(s, x + 0.25, y + 0.07, w - 0.35, 0.45, cabeza, size=size_cabeza, bold=True, color=color)
    texto(s, x + 0.25, y + 0.07 + size_cabeza / 40, w - 0.35, h - 0.5, cuerpo, size=size)


def dato(s, x, y, w, h, color, cifra, etiqueta, size_cifra=26):
    caja(s, x, y, w, h, GRIS_PANEL)
    caja(s, x, y, 0.08, h, color, forma=MSO_SHAPE.RECTANGLE)
    texto(s, x + 0.2, y + 0.04, w - 0.3, 0.55, cifra, size=size_cifra, bold=True, color=color)
    texto(s, x + 0.2, y + 0.58, w - 0.3, h - 0.6, etiqueta, size=13)


def mensaje(s, t, y=6.38, h=0.62, size=16):
    caja(s, 0.6, y, 12.13, h, AZUL)
    texto(s, 0.8, y, 11.8, h, t, size=size, color=BLANCO, anchor=MSO_ANCHOR.MIDDLE)


def flecha(s, x, y, w=0.6, h=0.4):
    caja(s, x, y, w, h, GRIS_FLECHA, forma=MSO_SHAPE.RIGHT_ARROW)


def codigo(s, x, y, w, h, lineas, color_titulo, cabeza):
    caja(s, x, y, w, h, TINTA)
    texto(s, x + 0.2, y + 0.1, w - 0.4, 0.4, cabeza, size=15, bold=True, color=color_titulo)
    texto(s, x + 0.2, y + 0.55, w - 0.4, h - 0.6, lineas, size=13, color=BLANCO, fuente="Consolas")


def imagen(s, nombre, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return s.shapes.add_picture(str(FIGURAS / nombre), Inches(x), Inches(y), **kw)


def notas(s, t):
    s.notes_slide.notes_text_frame.text = t


# %% PASO 2 · Preparar la plantilla: conservar portada y cierre, borrar el contenido antiguo
# ------------------------------------------------------------------------------
# La plantilla trae 18 diapositivas del proyecto anterior (datos simulados). Se conservan la 1
# (portada) y la 18 (cierre) y se eliminan las demás; las nuevas se insertan entre ambas.
# ------------------------------------------------------------------------------
prs = Presentation(str(PLANTILLA))
lista = prs.slides._sldIdLst
ids = list(lista)
for sld in ids[1:-1]:
    prs.part.drop_rel(sld.rId)
    lista.remove(sld)
BLANK = prs.slide_layouts[4]


def nueva():
    s = prs.slides.add_slide(BLANK)
    for ph in list(s.placeholders):            # quita los marcadores vacíos del pie
        ph._element.getparent().remove(ph._element)
    return s


# %% PASO 3 · Portada
# ------------------------------------------------------------------------------
portada = prs.slides[0]
for sh in portada.shapes:
    if not sh.has_text_frame:
        continue
    t = sh.text_frame.text
    nuevo = None
    if t.startswith("Auditoría híbrida"):
        nuevo = "Auditoría y prevención de fuga de datos en pipelines de Machine Learning generados por LLM"
    elif t.startswith("Detectar, medir"):
        nuevo = "Detectar, medir y reparar: un auditor determinista (AST + ejecución) y el Sesgo de Optimismo ΔOpt"
    elif t.startswith("Magíster"):
        nuevo = "Magíster en Data Science · Capstone Project I · Entrega final · Octubre 2026"
    if nuevo:
        p = sh.text_frame.paragraphs[0]
        p.runs[0].text = nuevo
        for r in p.runs[1:]:
            r.text = ""
        for extra in sh.text_frame.paragraphs[1:]:
            for r in extra.runs:
                r.text = ""
notas(portada, "(30 s) Presentarse. En una frase: los asistentes de IA ya escriben código de Machine Learning "
      "completo; este proyecto construye un auditor que encuentra un error silencioso de ese código —la fuga de "
      "datos—, mide cuánto infla la métrica y lo repara. El auditor es determinista: nunca usa un LLM para juzgar.")

# %% PASO 4 · Agenda
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Agenda")
items = [("El problema", "Fallos ruidosos vs. el fallo silencioso de la fuga"),
         ("Fundamentación", "Alternativas y decisiones de diseño"),
         ("Metodología", "Protocolo en 5 fases y arquitectura"),
         ("Resultados", "Eficacia del auditor y Sesgo de Optimismo ΔOpt"),
         ("Productos", "Dos Skills, plataforma web y uso responsable"),
         ("Cierre", "Interpretación, riesgos y Capstone II")]
for i, (cab, sub) in enumerate(items):
    x = 0.9 if i % 2 == 0 else 7.0
    y = 1.85 + (i // 2) * 1.5
    caja(s, x, y, 5.6, 1.2, GRIS_PANEL)
    o = caja(s, x + 0.2, y + 0.25, 0.7, 0.7, NARANJO, forma=MSO_SHAPE.OVAL)
    texto(s, x + 0.2, y + 0.25, 0.7, 0.7, str(i + 1), size=22, bold=True, color=BLANCO,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, x + 1.15, y + 0.15, 4.3, 0.45, cab, size=20, bold=True)
    texto(s, x + 1.15, y + 0.62, 4.3, 0.5, sub, size=14, color=GRIS_FLECHA)
notas(s, "(20 s) Recorrido: el problema y por qué solo la fuga; cómo se decidió el diseño; el protocolo; qué "
      "salió; los productos; y qué significa para Capstone II.")

# %% PASO 5 · El problema (R1)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "El problema: código que corre sin errores… y miente")
tarjeta(s, 0.6, 1.75, 3.9, 2.85, AZUL, "1. Los LLM ya programan ML",
        "Los asistentes de IA escriben pipelines completos: cargan datos, preprocesan, entrenan y reportan "
        "una métrica.", size=17)
tarjeta(s, 4.72, 1.75, 3.9, 2.85, AMBAR, "2. Fallos ruidosos",
        "Sintaxis, APIs inexistentes, dimensiones, memoria: el intérprete se detiene con un error. Nadie "
        "despliega un modelo que no corre.", size=17)
tarjeta(s, 8.84, 1.75, 3.9, 2.85, ROJO, "3. Fallo silencioso: la fuga",
        "El preprocesamiento «ve» las filas de prueba. El script corre de punta a punta y la métrica sale "
        "inflada: nadie la cuestiona.", size=17)
caja(s, 0.6, 4.8, 12.13, 1.35, GRIS_PANEL)
texto(s, 0.85, 4.88, 11.7, 1.2,
      [("Pregunta de investigación: ¿se puede detectar automáticamente la fuga de preprocesamiento en pipelines "
        "de scikit-learn, repararla sin cambiar el modelo y cuantificar cuánto optimismo introduce?",
        {"size": 18, "bold": True, "color": NARANJO})], anchor=MSO_ANCHOR.MIDDLE)
mensaje(s, "Evidencia de que importa: 294 artículos afectados por fuga en 17 campos científicos (Kapoor y Narayanan, 2023).")
notas(s, "(1 min) Observación R1 del comité: por qué solo la fuga. Hay siete fallos típicos del código generado; "
      "seis son ruidosos (excepción, linter o se detectan repitiendo la corrida). La fuga es el único que no "
      "interrumpe la ejecución y además mejora la métrica, por eso es el más peligroso: pasa a producción. "
      "Analogía para público no técnico: conocer las preguntas de la evaluación antes de rendirla; la nota sube, "
      "pero no mide lo aprendido.")

# %% PASO 6 · Cómo ocurre (R3)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "¿Cómo se fugan los datos? Misma operación, dos lugares")
codigo(s, 0.6, 1.7, 6.0, 2.55, ["X_esc = StandardScaler().fit_transform(X)",
                                "# el escalador vio TODAS las filas",
                                "cross_val_score(modelo, X_esc, y, cv=5)"], ROJO, "✗ Con fuga")
codigo(s, 6.73, 1.7, 6.0, 2.55, ["pipe = Pipeline([(\"esc\", StandardScaler()),",
                                 "                 (\"modelo\", modelo)])",
                                 "cross_val_score(pipe, X, y, cv=5)",
                                 "# en cada pliegue .fit() ve solo el entrenamiento"], VERDE, "✓ Sin fuga")
mecanismos = [(AZUL, "Escalado", "Media y varianza incluyen filas de prueba"),
              (AMBAR, "Imputación", "KNNImputer copia valores de vecinos de prueba"),
              (ROJO, "Selección", "SelectKBest mira las etiquetas de todas las filas"),
              (GRIS_FLECHA, "Combinada", "Los tres mecanismos a la vez")]
for i, (c, cab, cu) in enumerate(mecanismos):
    panel(s, 0.6 + i * 3.06, 4.45, 2.95, 1.7, c, cab, cu, size=13.5)
mensaje(s, "La regla: todo lo que «aprende» de los datos (.fit) debe ajustarse solo con el entrenamiento de cada pliegue.")
notas(s, "(1 min) Observación R3: mostrar el mecanismo. Arriba, el mismo escalador en dos lugares. Con fuga, la "
      "media y la desviación se calculan con las 569 filas antes de partir: cada pliegue de prueba ya fue visto. "
      "Sin fuga, el Pipeline reajusta en cada pliegue con el 80 % de entrenamiento. Abajo los cuatro mecanismos "
      "estudiados; la selección es la más grave porque usa la variable objetivo y. SMOTE no se midió porque no "
      "es parte de scikit-learn (desviación B2).")

# %% PASO 7 · Ejemplo visual (R3)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Ejemplo: el mismo pipeline con y sin fuga")
imagen(s, "fig_leakage_demo.png", 0.6, 1.75, w=8.0)
dato(s, 8.85, 1.75, 3.88, 1.1, AZUL, "ΔOpt = +0,0015", "Señal real (breast_cancer): casi invisible, efecto techo")
dato(s, 8.85, 3.0, 3.88, 1.1, ROJO, "ΔOpt = +0,048", "Señal nula: AUC 0,571 con fuga vs. 0,523 sin fuga")
dato(s, 8.85, 4.25, 3.88, 1.1, GRIS_FLECHA, "Verdad = 0,5", "Etiquetas permutadas: nada predice nada")
texto(s, 0.6, 5.3, 8.0, 0.8, "KNNImputer + StandardScaler + SelectKBest + MLP · mismos 5 pliegues (semilla 42) · "
      "leakage_demo.py", size=12, color=GRIS_FLECHA)
mensaje(s, "Con señal fuerte la fuga existe pero no se ve; con señal débil fabrica rendimiento que no existe.")
notas(s, "(1 min) Script leakage_demo.py, comentado paso a paso. Panel (a): con datos reales la diferencia es "
      "+0,0015 de AUC; la fuga existe pero el AUC ya está en 0,98, no hay margen. Panel (b): control con señal "
      "nula, 500 columnas de ruido y etiquetas permutadas; la verdad es 0,5. Con fuga el pipeline reporta 0,571: "
      "ese exceso es rendimiento fabricado. Pregunta probable: ¿por qué el eje Y completo? Para no exagerar "
      "visualmente diferencias pequeñas.")

# %% PASO 8 · Objetivos y criterios de éxito
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Objetivos verificables y criterios de éxito")
filas = [("OE1", AZUL, "Formalizar los mecanismos de fuga y distinguirlos de los fallos ruidosos",
          "96 pares con y sin fuga", "✔ Cumplido"),
         ("OE2", ROJO, "Auditor determinista (AST + ejecución) medido contra verdad conocida",
          "Recall independiente ≥ 0,80 y 0 FP", "◐ 9/9, IC95 [0,66; 1,00]"),
         ("OE3", AMBAR, "Cuantificar ΔOpt por tipo de fuga y 4 familias en 6 datasets",
          "H-esc.: |ΔOpt| < 0,001 en árboles", "✔ máx. 0,0002"),
         ("OE4", VERDE, "Dos Agent Skills y una plataforma web; reparación automática",
          "Reparación ≥ 90 % de scripts con fuga", "◐ 96/96 banco · 4/9 manual")]
texto(s, 1.75, 1.65, 6.0, 0.4, "Objetivo específico", size=14, bold=True, color=GRIS_FLECHA)
texto(s, 7.25, 1.65, 3.0, 0.4, "Criterio (umbral)", size=14, bold=True, color=GRIS_FLECHA)
texto(s, 10.15, 1.65, 2.6, 0.4, "Resultado", size=14, bold=True, color=GRIS_FLECHA)
for i, (oe, c, obj, crit, res) in enumerate(filas):
    y = 2.05 + i * 1.03
    caja(s, 0.6, y, 12.13, 0.9, GRIS_PANEL)
    caja(s, 0.6, y, 1.0, 0.9, c)
    texto(s, 0.6, y, 1.0, 0.9, oe, size=18, bold=True, color=BLANCO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, 1.75, y, 5.4, 0.9, obj, size=14.5, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, 7.25, y, 2.8, 0.9, crit, size=13.5, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, 10.15, y, 2.55, 0.9, res, size=14, bold=True, color=c, anchor=MSO_ANCHOR.MIDDLE)
mensaje(s, "Objetivo general: detectar, reparar y medir la fuga de datos en pipelines supervisados de scikit-learn.")
notas(s, "(1 min) Cada objetivo tiene un criterio con umbral y un resultado honesto. OE2: 9 de 9 casos con fuga "
      "y ningún falso positivo, pero con n = 9 el límite inferior del intervalo exacto es 0,66, así que no se "
      "puede afirmar el umbral de 0,80: por eso está a medias. OE4: la reparación funciona en los patrones "
      "conocidos (96/96) pero solo en 4 de 9 scripts escritos a mano; en los otros 5 declara «requiere revisión "
      "manual» y no toca el código. Los umbrales se fijarán antes de evaluar en Capstone II.")

# %% PASO 9 · Fundamentación: alternativas y decisiones
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Fundamentación: tres decisiones, tres alternativas descartadas")
tarjeta(s, 0.6, 1.75, 3.9, 3.0, AZUL, "Detector",
        [("Elegido: AST + ejecución (determinista)", {"bold": True, "size": 16.5}),
         ("Descartado: LLM como juez — prefiere sus propias respuestas y varía entre corridas (Zheng et al., 2023).",
          {"size": 16})])
tarjeta(s, 4.72, 1.75, 3.9, 3.0, ROJO, "Métrica",
        [("Elegido: ΔOpt pareado, mismos pliegues", {"bold": True, "size": 16.5}),
         ("Descartado: brecha de generalización — mide sobreajuste, que existe aun sin fuga.", {"size": 16})])
tarjeta(s, 8.84, 1.75, 3.9, 3.0, AMBAR, "Verdad de referencia",
        [("Elegido: oráculo programático + 13 casos manuales", {"bold": True, "size": 16.5}),
         ("Descartado: anotadores humanos — subjetivos (y el comité pidió evitarlos).", {"size": 16})])
caja(s, 0.6, 4.95, 12.13, 1.2, GRIS_PANEL)
texto(s, 0.85, 4.98, 11.7, 1.15,
      [("Brecha: ", {"bold": True, "color": NARANJO, "size": 17}),
       ("los detectores existentes (Yang et al., 2022; Drobnjaković et al., 2024; AlOmar et al., 2025) son "
        "estáticos, se validaron con código humano y no miden el daño ni reparan.", {"size": 16})],
      anchor=MSO_ANCHOR.MIDDLE)
mensaje(s, "Un instrumento de medición debe dar el mismo veredicto en cada corrida y no tener interés en el resultado.")
notas(s, "(1,5 min) Criterio 2: argumentar con la literatura, no describirla. Detector: el análisis estático es "
      "reproducible pero ciego a valores y alias; la ejecución ve el ajuste real; un LLM juez no sirve como "
      "instrumento. Métrica: la validación cruzada anidada corrige la selección de modelos (Varma y Simon, 2006; "
      "Cawley y Talbot, 2010) pero no aísla un error concreto del código; la comparación pareada sí. Verdad: el "
      "oráculo conoce la etiqueta con certeza, pero hay circularidad (mismo autor del generador y del auditor), "
      "por eso los 13 casos manuales. Ambroise y McLachlan (2002) justifican agregar columnas de ruido.")

# %% PASO 10 · Métrica ΔOpt (R9)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "La métrica: Sesgo de Optimismo Inducido por Fuga (ΔOpt)")
for i, (cab, val, sub, c) in enumerate([("Pipeline CON fuga", "M_leaky", "preprocesamiento con todas las filas", ROJO),
                                         ("Pipeline SIN fuga", "M_clean", "preprocesamiento dentro de Pipeline", VERDE),
                                         ("Sesgo de optimismo", "ΔOpt", "= M_leaky − M_clean", NARANJO)]):
    x = 0.9 + i * 4.15
    caja(s, x, 1.85, 3.3, 2.2, BLANCO, GRIS_BORDE)
    texto(s, x, 1.95, 3.3, 0.45, cab, size=16, bold=True, color=TINTA, align=PP_ALIGN.CENTER)
    texto(s, x, 2.45, 3.3, 0.8, val, size=36, bold=True, color=c, align=PP_ALIGN.CENTER)
    texto(s, x, 3.3, 3.3, 0.6, sub, size=13.5, color=GRIS_FLECHA, align=PP_ALIGN.CENTER)
    if i < 2:
        flecha(s, x + 3.4, 2.75, 0.6, 0.45)
caja(s, 0.9, 4.35, 11.5, 1.75, GRIS_PANEL)
texto(s, 1.1, 4.42, 11.4, 1.65,
      ["Mismo dataset · mismos 5 pliegues (semilla 42) · mismo modelo · mismos hiperparámetros.",
       "Solo cambia DÓNDE se ajusta el preprocesamiento → la diferencia es causalmente atribuible a la fuga.",
       "Sin fuga, ΔOpt = 0 por identidad. ROC-AUC en clasificación, R² en regresión."], size=16)
mensaje(s, "Reemplaza a la «brecha de generalización auditada», que se confundía con el sobreajuste (observación R9).")
notas(s, "(1 min) Observación R9. ΔOpt es una intervención pareada: dos versiones idénticas salvo un detalle. "
      "Unidades: puntos de la métrica; +0,05 de AUC significa que la fuga infló el AUC en 0,05. Pregunta "
      "probable: ¿por qué no la brecha entrenamiento-prueba? Porque un árbol sin poda tiene brecha enorme sin "
      "ninguna fuga; mezcla dos fenómenos.")

# %% PASO 11 · Metodología: 5 fases (R4, R5, R6)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Metodología: protocolo en cinco fases deterministas")
fases = [("1", "Banco de control", "96 pares con y sin fuga (oráculo programático)", AZUL),
         ("2", "Auditoría", "AST + verificación en ejecución", ROJO),
         ("3", "Eficacia", "Precisión, recall, VPN vs. verdad y línea base", AMBAR),
         ("4", "Sesgo ΔOpt", "Por par, pruebas por familia, control positivo", VERDE),
         ("5", "Empaquetado", "2 Skills + plataforma web", NARANJO)]
for i, (n, cab, cu, c) in enumerate(fases):
    x = 0.6 + i * 2.47
    caja(s, x, 1.75, 2.25, 2.3, BLANCO, GRIS_BORDE)
    caja(s, x, 1.75, 2.25, 0.09, c, forma=MSO_SHAPE.RECTANGLE)
    texto(s, x + 0.1, 1.9, 2.05, 0.5, f"Fase {n}", size=14, bold=True, color=c)
    texto(s, x + 0.1, 2.3, 2.05, 0.5, cab, size=17, bold=True)
    texto(s, x + 0.1, 2.85, 2.05, 1.2, cu, size=13)
for i, (cif, et, c) in enumerate([("6 datasets", "nativos de sklearn: 3 clasificación + 3 regresión", AZUL),
                                  ("4 familias", "RandomForest · HistGB · Lineal · MLP", ROJO),
                                  ("4 fugas", "escalado · imputación · selección · combinada", AMBAR),
                                  ("192 scripts", "= 6 × 4 × 4 pares × 2 versiones", VERDE)]):
    dato(s, 0.6 + i * 3.06, 4.3, 2.95, 1.25, c, cif, et, size_cifra=24)
texto(s, 0.6, 5.65, 12.1, 0.5, "Fase 0: EDA de calidad (0 duplicados, atípicos ≤ 1,24 %), 10 % de faltantes simulados "
      "y 50 columnas de ruido. Todo con semilla 42 y un comando: python run_all.py",
      size=13, color=GRIS_FLECHA)
mensaje(s, "Sin anotadores humanos: la etiqueta «con fuga / sin fuga» se conoce por construcción (observación R5).")
notas(s, "(1,5 min) Observaciones R4, R5 y R6. Las 19 etapas del diseño anterior se redujeron a 5, cada una con un "
      "script y un archivo de resultados (Tabla 3.1 del informe). Clasificación y regresión supervisadas; "
      "clustering queda fuera porque no hay objetivo contra el cual medir optimismo. california_housing se "
      "submuestreó a 3.000 filas por el costo de KNNImputer (desviación B1). Las columnas de ruido existen porque "
      "la fuga por selección solo aparece si hay variables inútiles entre las cuales elegir.")

# %% PASO 12 · Arquitectura (R2)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Arquitectura: quién genera y quién audita")
caja(s, 0.6, 2.6, 2.2, 1.6, TINTA)
texto(s, 0.6, 2.6, 2.2, 1.6, ["Script de ML", ("(banco, manual o LLM)", {"size": 13, "bold": False})],
      size=17, bold=True, color=BLANCO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
flecha(s, 2.9, 2.0)
flecha(s, 2.9, 4.4)
tarjeta(s, 3.6, 1.6, 3.9, 1.95, AZUL, "AST (estático)",
        "Lee sin ejecutar. Reglas R-Split, R-CV y R-Test sobre el orden de .fit y la partición.", size=13.5)
tarjeta(s, 3.6, 3.85, 3.9, 2.05, NARANJO, "Ejecución (dinámico)",
        "Cuenta con cuántas filas se ajusta cada transformador: si son todas, vio la prueba.", size=13.5)
flecha(s, 7.65, 3.15, 0.6, 0.5)
panel(s, 8.4, 1.6, 4.33, 1.3, ROJO, "FUGA", "Basta que UNO de los dos la detecte (híbrido)")
panel(s, 8.4, 3.07, 4.33, 1.3, VERDE, "0 FUGAS", "Sin evidencia de las 3 reglas — no «sin ninguna fuga»")
panel(s, 8.4, 4.54, 4.33, 1.3, AMBAR, "REPARACIÓN", "Encierra en Pipeline sin cambiar modelo ni pliegues")
mensaje(s, "Los LLM son solo sujetos de estudio (Capstone II). Ningún LLM participa en la medición.")
notas(s, "(1 min) Observación R2: separar generador y auditor evita el sesgo de autoconfirmación. Analogía: el AST "
      "lee la receta; la ejecución mira al cocinero mientras cocina. Cada uno ve cosas que el otro no: el AST no "
      "ve alias de importación; la ejecución no ve ajustes parciales (70 % de las filas). El veredicto limpio "
      "siempre se acompaña del aviso de alcance.")

# %% PASO 13 · Resultados del auditor
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Resultados: ¿qué tan bien detecta el auditor?")
dato(s, 0.6, 1.75, 3.9, 1.15, GRIS_FLECHA, "192 / 192", "Banco de control — evaluación CIRCULAR, no prueba generalización")
dato(s, 4.72, 1.75, 3.9, 1.15, VERDE, "9 / 9", "Híbrido en 13 scripts manuales · 0 falsos positivos")
dato(s, 8.84, 1.75, 3.9, 1.15, AMBAR, "[0,66; 1,00]", "IC95 % exacto del recall híbrido (n = 9)")
encab = ["Detector (13 manuales)", "Detecta", "Falsos +", "Recall [IC95 %]"]
filas = [("Regla simple (línea base)", "8/9", "1", "0,89 [0,52; 1,00]"),
         ("Solo AST", "8/9", "0", "0,89 [0,52; 1,00]"),
         ("Solo ejecución", "7/9", "0", "0,78 [0,40; 0,97]"),
         ("Híbrido", "9/9", "0", "1,00 [0,66; 1,00]")]
xs, ws = [0.6, 5.1, 7.0, 8.9], [4.5, 1.9, 1.9, 3.83]
caja(s, 0.6, 3.1, 12.13, 0.5, TINTA, forma=MSO_SHAPE.RECTANGLE)
for x, w, e in zip(xs, ws, encab):
    texto(s, x + 0.15, 3.1, w - 0.2, 0.5, e, size=15, bold=True, color=BLANCO, anchor=MSO_ANCHOR.MIDDLE)
for j, fila in enumerate(filas):
    y = 3.6 + j * 0.62
    caja(s, 0.6, y, 12.13, 0.62, GRIS_PANEL if j % 2 == 0 else BLANCO, forma=MSO_SHAPE.RECTANGLE)
    for x, w, v in zip(xs, ws, fila):
        texto(s, x + 0.15, y, w - 0.2, 0.62, v, size=15, bold=(j == 3), color=VERDE if j == 3 else TINTA,
              anchor=MSO_ANCHOR.MIDDLE)
mensaje(s, "El AST falla con el alias «as SS»; la ejecución, con ajustes parciales: fallan en casos distintos → híbrido.")
notas(s, "(1,5 min) Primero la honestidad: el 100 % del banco es circular; el generador y el auditor tienen el "
      "mismo autor. La prueba real son 13 scripts escritos a mano con estilos que el generador no produce. La "
      "regla simple (grep: .fit sin Pipeline) rinde igual que el AST en recall pero con un falso positivo y "
      "pierde 24 de 96 fugas en el banco. Los intervalos se solapan: la ventaja del híbrido es una señal que "
      "Capstone II debe confirmar con 60 o más scripts, no una mejora demostrada. Clopper-Pearson porque es "
      "exacto para muestras pequeñas.")

# %% PASO 14 · Resultados ΔOpt
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Resultados: ¿cuánto infla la fuga la métrica?")
imagen(s, "fig_delta_opt_heatmap.png", 1.85, 1.55, w=9.6)
dato(s, 0.6, 4.85, 2.9, 1.3, AZUL, "≈ 0", "Clasificación: |ΔOpt| ≤ 0,0034 (AUC limpio mediano 0,991)", size_cifra=24)
dato(s, 3.68, 4.85, 2.9, 1.3, ROJO, "+0,074", "Máximo en regresión: selección, california, MLP", size_cifra=24)
dato(s, 6.76, 4.85, 2.9, 1.3, VERDE, "0,0002", "Escalado en árboles: confirma la teoría", size_cifra=24)
dato(s, 9.84, 4.85, 2.89, 1.3, AMBAR, "p Holm 0,297", "MLP en regresión (p = 0,037 sin corregir)", size_cifra=22)
mensaje(s, "El efecto es pequeño con datos «fáciles»; la vulnerabilidad de la MLP es exploratoria, no confirmada.")
notas(s, "(1,5 min) Tres hechos. (1) Clasificación: ΔOpt prácticamente nulo por efecto techo. (2) Regresión: "
      "aparece optimismo, concentrado en la selección con california_housing (R² 0,615 → 0,689 en MLP). (3) "
      "Escalado en árboles: máximo 0,0002, porque una transformación monótona no cambia los cortes de un árbol. "
      "La MLP en regresión tiene el mayor ΔOpt medio (0,0169; IC bootstrap [0,0027; 0,0346]), pero con 8 pruebas "
      "el ajuste de Holm sube su p a 0,297: se reporta como exploratorio. Un par degenerado (MLP divergió) se "
      "excluyó con una regla definida después de verlo; se declara.")

# %% PASO 15 · Control positivo
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Control positivo: cuando la señal es débil, la fuga sí daña")
imagen(s, "fig_noise_sensitivity.png", 0.6, 1.6, h=4.6)
dato(s, 9.3, 1.75, 3.43, 1.25, ROJO, "+0,20 AUC", "Señal nula, breast_cancer, lineal, 1.000 columnas de ruido")
dato(s, 9.3, 3.15, 3.43, 1.25, ROJO, "+0,26 R²", "Señal nula, diabetes, lineal")
dato(s, 9.3, 4.55, 3.43, 1.25, AZUL, "Plano", "Con etiquetas reales ΔOpt no cambia con el ruido")
mensaje(s, "La gravedad depende de la razón señal/ruido, no es una propiedad fija de cada fuga o familia.")
notas(s, "(1 min) Como la selección casi no inflaba el AUC, había que comprobar que ΔOpt detecta la fuga cuando "
      "importa. Fila superior: etiquetas reales, curvas planas: con 8 variables muy informativas el selector "
      "nunca elige ruido. Fila inferior: etiquetas permutadas; la verdad es AUC 0,5, y ΔOpt crece con el ruido "
      "hasta +0,20. Implicancia: para Capstone II se reformula OE3 como ΔOpt condicionado a la señal/ruido.")

# %% PASO 16 · Productos (R7, R8) y uso responsable
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Productos: tres vías, un mismo auditor")
tarjeta(s, 0.6, 1.75, 3.9, 2.7, AZUL, "Skill audit-ml-leakage",
        "Para agentes de programación: audita y repara un script existente; re-audita hasta «0 fugas».", size=17)
tarjeta(s, 4.72, 1.75, 3.9, 2.7, VERDE, "Skill create-leakage-free",
        "Reglas y plantilla para crear pipelines sin fuga desde el inicio, con autoverificación.", size=17)
tarjeta(s, 8.84, 1.75, 3.9, 2.7, NARANJO, "Plataforma web (Streamlit)",
        "Pegar o subir un .py, ver líneas en rojo, descargar el código corregido y explorar ΔOpt.", size=17)
panel(s, 0.6, 4.65, 6.0, 1.5, ROJO, "Riesgos y medidas",
      "Falsa seguridad → aviso de alcance visible · Código malicioso → la web nunca ejecuta, límite 200.000 "
      "caracteres, HTML escapado", size=14.5)
panel(s, 6.73, 4.65, 6.0, 1.5, AZUL, "Reproducible y auditable",
      "Repositorio público en GitHub · 40 pruebas automáticas · CLI con código de salida para CI · "
      "dos corridas producen archivos idénticos", size=14.5)
mensaje(s, "Repositorio: github.com/witman92/ml-leakage-auditor-capstone  ·  streamlit run 06_web_app/app.py")
notas(s, "(1,5 min) Observaciones R7 y R8. Las Skills sirven a quien trabaja con un agente; la web a quien no. "
      "Comparten el mismo núcleo (03_auditor), así que dan el mismo veredicto. Si hay tiempo, demo en vivo de "
      "la web con el ejemplo 03 (SelectKBest antes de la validación cruzada): marca la línea, repara y muestra "
      "el diff. Ética: datos públicos y anonimizados de scikit-learn sin fines clínicos; el código de LLM se "
      "ejecutará solo en contenedor sin red. Falta una prueba de uso con personas (Capstone II).")

# %% PASO 17 · Interpretación, riesgos y ajustes (criterio 5)
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Interpretación: qué significa para el proyecto")
caja(s, 0.6, 1.75, 6.0, 4.4, BLANCO, GRIS_BORDE)
caja(s, 0.6, 1.75, 6.0, 0.09, AZUL, forma=MSO_SHAPE.RECTANGLE)
texto(s, 0.75, 1.93, 5.7, 0.5, "Lo que la evidencia sostiene", size=18, bold=True, color=AZUL)
texto(s, 0.75, 2.5, 5.7, 3.6,
      ["• El híbrido es más robusto que cada parte; no que sea «casi perfecto».",
       "• La fuga fue pequeña porque los datos son fáciles, no porque sea inocua.",
       "• Factible como instrumento (detectar, reparar, medir).",
       "• Débil como afirmación sobre «qué familia sufre más»: 3 datasets por tarea."], size=17)
caja(s, 6.73, 1.75, 6.0, 4.4, BLANCO, GRIS_BORDE)
caja(s, 6.73, 1.75, 6.0, 0.09, ROJO, forma=MSO_SHAPE.RECTANGLE)
texto(s, 6.88, 1.93, 5.7, 0.5, "Riesgos para Capstone II → ajuste", size=18, bold=True, color=ROJO)
texto(s, 6.88, 2.5, 5.7, 3.6,
      ["• Código de LLM más variado → ampliar la prueba independiente a ≥ 60 scripts.",
       "• Alias y funciones auxiliares → mejorar el AST y el refactorizador.",
       "• Efecto techo → datasets de señal débil; ΔOpt por tramo de señal/ruido.",
       "• Ejecutar código no confiable → contenedor sin red ni credenciales."], size=17)
mensaje(s, "OE3 se reformula: estimar ΔOpt condicionado a la razón señal/ruido, no un número único por familia.")
notas(s, "(1,5 min) Este es el criterio 5: decir con palabras propias qué significan los resultados y qué se ajusta. "
      "Ninguna afirmación va más allá de la evidencia: el 9/9 tiene un intervalo amplio, la MLP no sobrevive a "
      "Holm, el banco es circular. Limitaciones de datos: seis datasets clásicos que los LLM probablemente "
      "conocen, submuestreo de california y efecto techo.")

# %% PASO 18 · Conclusiones y próximos pasos
# ------------------------------------------------------------------------------
s = nueva()
titulo(s, "Conclusiones y próximo paso")
concl = [(VERDE, "✔", "Detectar es posible: el diseño híbrido encontró 9/9 fugas en código escrito a mano."),
         (VERDE, "✔", "Reparar sin cambiar el modelo funciona en los patrones habituales (96/96; 4/9 manual)."),
         (VERDE, "✔", "ΔOpt cuantifica la fuga y la separa del sobreajuste; su magnitud depende de la señal/ruido."),
         (AMBAR, "!", "Falta lo esencial: medir la fuga en código de LLM reales (Capstone II).")]
for i, (c, ic, t) in enumerate(concl):
    y = 1.75 + i * 0.85
    o = caja(s, 0.6, y, 0.6, 0.6, c, forma=MSO_SHAPE.OVAL)
    texto(s, 0.6, y, 0.6, 0.6, ic, size=18, bold=True, color=BLANCO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    texto(s, 1.4, y, 11.3, 0.6, t, size=18, anchor=MSO_ANCHOR.MIDDLE)
texto(s, 0.6, 5.2, 12.1, 0.45, "Hoja de ruta · Capstone II", size=18, bold=True, color=NARANJO)
for i, t in enumerate(["6–8 LLM reales", "Contenedor aislado", "≥ 60 casos manuales", "Señal débil", "Piloto con usuarios"]):
    caja(s, 0.6 + i * 2.45, 5.7, 2.3, 0.7, AZUL)
    texto(s, 0.6 + i * 2.45, 5.7, 2.3, 0.7, t, size=15, bold=True, color=BLANCO, align=PP_ALIGN.CENTER,
          anchor=MSO_ANCHOR.MIDDLE)
notas(s, "(1 min) Responder la pregunta de investigación en tres partes: detectar sí, reparar parcialmente, "
      "cuantificar sí con matices. Cerrar: el instrumento está listo y es reproducible; el siguiente ciclo lo "
      "aplica a código de LLM reales con umbrales fijados de antemano.")

# %% PASO 19 · Cierre: mover la diapositiva de agradecimiento al final y guardar
# ------------------------------------------------------------------------------
cierre = lista[1]                      # tras borrar, la segunda entrada es la diapositiva de cierre
lista.remove(cierre)
lista.append(cierre)
notas(prs.slides[len(prs.slides) - 1],
      "Agradecer y abrir preguntas. Preguntas probables: (1) ¿Por qué el 100 % del banco no vale? Circularidad. "
      "(2) ¿Por qué no un LLM como juez? Reproducibilidad y sesgo de autopreferencia. (3) ¿Qué no detecta? "
      "Fugas temporales, de dominio, duplicados, alias, exec. (4) ¿Por qué ΔOpt ≈ 0 en clasificación? Efecto "
      "techo y señal fuerte; el control positivo muestra que sí mide cuando importa. (5) ¿Qué es Holm? Ajuste "
      "escalonado de los valores p cuando se hacen varias pruebas a la vez.")
prs.save(str(SALIDA))
print(f"{len(prs.slides)} diapositivas → {SALIDA}")
